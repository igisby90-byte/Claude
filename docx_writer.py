"""수집한 기사를 메일 양식 그대로 워드(.docx) 파일로 작성한다."""

from __future__ import annotations

from datetime import date

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor


def add_hyperlink(paragraph, url: str, text: str,
                  color: str = "0563C1", underline: bool = True):
    """문단에 클릭 가능한 하이퍼링크 run 을 추가한다.

    python-docx 는 하이퍼링크 헬퍼를 기본 제공하지 않아 직접 XML 을 구성한다.
    """
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")

    if color:
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        rpr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)

    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)
    return hyperlink


def build_document(config: dict, results: dict, week_label: str,
                   start_date: date, end_date: date):
    """전체 보고서 문서를 생성해 Document 객체를 반환한다.

    results 구조:
        { fund_name: { asset_label: [Article, ...], ... }, ... }
    """
    doc = docx.Document()

    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10)

    # 제목
    title = doc.add_paragraph()
    run = title.add_run(f"[이지스자산운용] {config['report_title']} ({week_label})")
    run.bold = True
    run.font.size = Pt(13)

    period = doc.add_paragraph()
    period.add_run(
        f"수집 기간: {start_date:%Y-%m-%d}(월) ~ {end_date:%Y-%m-%d}(일)"
    ).italic = True

    doc.add_paragraph("")

    # 인사말
    for line in config.get("intro_lines", []):
        doc.add_paragraph(line)
    doc.add_paragraph("")

    # 펀드별 본문
    for fund_cfg in config["funds"]:
        fund_name = fund_cfg["fund"]
        fund_para = doc.add_paragraph()
        fund_para.add_run(f"[{fund_name} 투자자산 관련 기사]").bold = True

        for idx, asset in enumerate(fund_cfg["assets"], start=1):
            label = asset["label"]
            asset_para = doc.add_paragraph()
            asset_para.add_run(f"{idx}. {label}").bold = True

            articles = results.get(fund_name, {}).get(label, [])
            if not articles:
                p = doc.add_paragraph("    - 전 주 관련 기사 없음", style="Normal")
                p.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)
                continue

            for art in articles:
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Pt(18)
                p.add_run("- ")
                add_hyperlink(p, art.link, art.title)
                p.add_run(f"  ({art.pub_date:%m-%d})").font.size = Pt(8)

        doc.add_paragraph("")

    # 맺음말
    for line in config.get("closing_lines", []):
        doc.add_paragraph(line)

    return doc
