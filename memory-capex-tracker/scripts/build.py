#!/usr/bin/env python3
"""data/*.csv → tracker.xlsx + dashboard.html 동시 생성.

사용법:  python scripts/build.py
데이터를 고친 뒤 이 스크립트만 다시 돌리면 엑셀과 HTML이 함께 갱신됩니다.
"""
import csv
import json
import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.marker import Marker
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

# 팔레트 (dataviz 검증 통과: 4슬롯 categorical)
C = {"msft": "2A78D6", "googl": "008300", "meta": "E87BA4", "amzn": "EDA100",
     "nominal": "2A78D6", "real": "EDA100", "dfii": "4A3AA7",
     "ink": "0B0B0B", "muted": "898781", "grid": "E1E0D9"}


def read_csv(name):
    with (DATA / name).open() as f:
        return list(csv.DictReader(f))


def load():
    params = {r["key"]: r["value"] for r in read_csv("params.csv")}
    capex = [{"q": r["quarter"], "msft": float(r["msft"]), "googl": float(r["googl"]),
              "meta": float(r["meta"]), "amzn": float(r["amzn"]),
              "status": r["status"], "note": r.get("note", "")} for r in read_csv("capex.csv")]
    dram = [{"q": r["quarter"], "low": float(r["qoq_low_pct"]), "mid": float(r["qoq_mid_pct"]),
             "high": float(r["qoq_high_pct"]), "status": r["status"]} for r in read_csv("dram.csv")]
    dfii = [{"m": r["month"], "v": float(r["dfii10"]) if r["dfii10"] else None,
             "status": r["status"]} for r in read_csv("dfii10.csv")]
    return params, capex, dram, dfii


# ---------------------------------------------------------------- HTML
def build_html(params, capex, dram, dfii):
    payload = {
        "generated": datetime.date.today().isoformat(),
        "params": {"memory_share_pct": float(params["memory_share_pct"])},
        "capex": capex, "dram": dram, "dfii10": dfii,
    }
    tpl = (ROOT / "scripts" / "template.html").read_text()
    out = tpl.replace("/*__DATA__*/", json.dumps(payload, ensure_ascii=False))
    (ROOT / "dashboard.html").write_text(out)
    print(f"wrote {ROOT/'dashboard.html'}")


# ---------------------------------------------------------------- Excel
def style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color=C["ink"])
        cell.fill = PatternFill("solid", fgColor="F0EFEC")
        cell.alignment = Alignment(horizontal="center")


def line_style(series, hex_color, width_pt=1.75, dashed=False):
    lp = LineProperties(solidFill=hex_color, w=int(width_pt * 12700))
    if dashed:
        lp.prstDash = "dash"
    series.graphicalProperties.line = lp
    series.smooth = False


def build_xlsx(params, capex, dram, dfii):
    wb = Workbook()
    n_cap = len(capex)
    n_dram = len(dram)

    # --- Params
    ws = wb.active
    ws.title = "Params"
    ws.append(["항목", "값", "메모"])
    ws.append(["메모리·부품 비중 (%)", float(params["memory_share_pct"]),
               "MSFT 자백치 $25B/$190B=13.2% — 이 셀만 바꾸면 Derived 전체 재계산"])
    ws.append(["보수적 하단 (%)", float(params["memory_share_low_pct"]), "Meta 상향폭 기준 약 7~8%"])
    ws.append(["DRAM 지수 기준분기", params["index_base"], "=100"])
    style_header(ws, 3)
    for col, w in zip("ABC", (26, 10, 70)):
        ws.column_dimensions[col].width = w
    ws["B2"].font = Font(bold=True, size=13, color=C["nominal"])

    # --- Capex
    ws = wb.create_sheet("Capex")
    ws.append(["분기", "MSFT", "GOOGL", "META", "AMZN", "합계($B)", "상태", "메모"])
    for i, r in enumerate(capex, start=2):
        ws.append([r["q"], r["msft"], r["googl"], r["meta"], r["amzn"],
                   f"=SUM(B{i}:E{i})", r["status"], r["note"]])
    style_header(ws, 8)
    for col, w in zip("ABCDEFGH", (9, 8, 8, 8, 8, 10, 10, 46)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=6):
        for c in row:
            c.number_format = "#,##0.0"

    # --- DRAM
    ws = wb.create_sheet("DRAM")
    ws.append(["분기", "QoQ 하단%", "QoQ 중간%", "QoQ 상단%", "지수(중간)", "지수(하단)", "지수(상단)", "상태"])
    for i, r in enumerate(dram, start=2):
        prev = "100" if i == 2 else f"E{i-1}"
        prev_l = "100" if i == 2 else f"F{i-1}"
        prev_h = "100" if i == 2 else f"G{i-1}"
        ws.append([r["q"], r["low"], r["mid"], r["high"],
                   f"={prev}*(1+C{i}/100)", f"={prev_l}*(1+B{i}/100)", f"={prev_h}*(1+D{i}/100)",
                   r["status"]])
    style_header(ws, 8)
    for col, w in zip("ABCDEFGH", (9, 10, 10, 10, 11, 11, 11, 22)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=7):
        for c in row:
            c.number_format = "#,##0"

    # --- Derived (수식: Params!B2 비중을 바꾸면 전체 재계산)
    ws = wb.create_sheet("Derived")
    ws.append(["분기", "총 capex($B)", "메모리 지출($B)", "DRAM 지수", "실질 메모리($B)",
               "실질 총 capex($B)", "명목 지출 지수", "실질 구매량 지수", "명목 YoY%", "실질 YoY%", "착시 폭%p"])
    dram_row = {r["q"]: i for i, r in enumerate(dram, start=2)}
    cap_row = {r["q"]: i for i, r in enumerate(capex, start=2)}
    base_q = params["index_base"]
    for i, r in enumerate(capex, start=2):
        q = r["q"]
        dr = dram_row.get(q)
        ws.cell(row=i, column=1, value=q)
        ws.cell(row=i, column=2, value=f"=Capex!F{cap_row[q]}")
        ws.cell(row=i, column=3, value=f"=B{i}*Params!$B$2/100")
        if dr:
            bq = cap_row[base_q]
            ws.cell(row=i, column=4, value=f"=DRAM!E{dr}")
            ws.cell(row=i, column=5, value=f"=C{i}/(D{i}/100)")
            ws.cell(row=i, column=6, value=f"=B{i}-C{i}+E{i}")
            ws.cell(row=i, column=7, value=f"=C{i}/C${bq}*100")
            ws.cell(row=i, column=8, value=f"=E{i}/E${bq}*100")
        prev = cap_row.get(str(int(q[:4]) - 1) + q[4:])
        if prev:
            ws.cell(row=i, column=9, value=f"=(B{i}/B{prev}-1)*100")
            ws.cell(row=i, column=10, value=f"=(IF(F{i}=\"\",B{i},F{i})/IF(F{prev}=\"\",B{prev},F{prev})-1)*100")
            ws.cell(row=i, column=11, value=f"=I{i}-J{i}")
    style_header(ws, 11)
    for idx in range(1, 12):
        ws.column_dimensions[get_column_letter(idx)].width = 13
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=11):
        for c in row:
            c.number_format = "#,##0.0"

    # --- DFII10
    ws = wb.create_sheet("DFII10")
    ws.append(["월", "DFII10 (%)", "상태"])
    for r in dfii:
        ws.append([r["m"], r["v"], r["status"]])
    style_header(ws, 3)
    for col, w in zip("ABC", (12, 12, 12)):
        ws.column_dimensions[col].width = w

    # ---------------- Charts sheet
    cs = wb.create_sheet("Charts")
    cs["A1"] = "메모리 착시 트래커 — Params!B2(비중)를 바꾸면 Derived 기반 차트가 재계산됩니다"
    cs["A1"].font = Font(bold=True, size=12)

    # 1) 4사 capex 누적 막대
    ch = BarChart()
    ch.type, ch.grouping, ch.overlap = "col", "stacked", 100
    ch.title = "빅테크 4사 분기 capex (명목, $B)"
    ch.height, ch.width = 9, 17
    data = Reference(wb["Capex"], min_col=2, max_col=5, min_row=1, max_row=n_cap + 1)
    cats = Reference(wb["Capex"], min_col=1, min_row=2, max_row=n_cap + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    for s, key in zip(ch.series, ("msft", "googl", "meta", "amzn")):
        s.graphicalProperties.solidFill = C[key]
        s.graphicalProperties.line.noFill = True
    ch.gapWidth = 40
    cs.add_chart(ch, "A3")

    # 2) DRAM 지수
    ch = LineChart()
    ch.title = "서버 DRAM 계약가 지수 (2025Q1=100, 하단·중간·상단)"
    ch.height, ch.width = 9, 17
    data = Reference(wb["DRAM"], min_col=5, max_col=7, min_row=1, max_row=n_dram + 1)
    cats = Reference(wb["DRAM"], min_col=1, min_row=2, max_row=n_dram + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    line_style(ch.series[0], C["nominal"])
    line_style(ch.series[1], C["muted"], 1.0, dashed=True)
    line_style(ch.series[2], C["muted"], 1.0, dashed=True)
    for s in ch.series:
        s.marker = Marker(symbol="circle", size=6)
    cs.add_chart(ch, "J3")

    # 3) 명목 지출 vs 실질 구매량 지수
    ch = LineChart()
    ch.title = "메모리: 명목 지출 vs 실질 구매량 (지수, 2025Q1=100)"
    ch.height, ch.width = 9, 17
    data = Reference(wb["Derived"], min_col=7, max_col=8, min_row=1, max_row=n_cap + 1)
    cats = Reference(wb["Derived"], min_col=1, min_row=2, max_row=n_cap + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    line_style(ch.series[0], C["nominal"])
    line_style(ch.series[1], C["real"])
    for s in ch.series:
        s.marker = Marker(symbol="circle", size=6)
    cs.add_chart(ch, "A22")

    # 4) 명목 vs 실질 YoY
    ch = LineChart()
    ch.title = "총 capex YoY: 명목 vs 실질 (%) — 간격이 착시 폭"
    ch.height, ch.width = 9, 17
    data = Reference(wb["Derived"], min_col=9, max_col=10, min_row=1, max_row=n_cap + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    line_style(ch.series[0], C["nominal"])
    line_style(ch.series[1], C["real"])
    for s in ch.series:
        s.marker = Marker(symbol="circle", size=6)
    cs.add_chart(ch, "J22")

    # 5) DFII10
    ch = LineChart()
    ch.title = "DFII10 — 10년 TIPS 실질금리 (%, 월말)"
    ch.height, ch.width = 9, 34
    n_df = len(dfii)
    data = Reference(wb["DFII10"], min_col=2, min_row=1, max_row=n_df + 1)
    cats = Reference(wb["DFII10"], min_col=1, min_row=2, max_row=n_df + 1)
    ch.add_data(data, titles_from_data=True)
    ch.set_categories(cats)
    line_style(ch.series[0], C["dfii"])
    ch.series[0].marker = Marker(symbol="circle", size=5)
    cs.add_chart(ch, "A41")

    out = ROOT / "tracker.xlsx"
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    args = load()
    build_html(*args)
    build_xlsx(*args)
